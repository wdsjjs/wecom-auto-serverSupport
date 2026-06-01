from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .db import connect, ensure_schema
from .textutil import first_non_empty, json_dumps, normalize_text, stable_id


PRODUCT_SHEET = "5 产品常规信息"
SHIPPING_SHEET = "1 产品发货状态"
RECOMMENDATION_SHEET = "10 补剂推荐"
REPORT_SHEET = "13 产品检测报告"
NOTICE_SHEET = "3 限时通知"
GENERIC_SHEET_TYPES = {
    "2 促单活动": "promotion_notice",
    "6 论文表": "research_evidence",
    "7 L0级注意事项": "safety_policy",
    "12原料专利、认证等材料": "raw_material_certification",
    "15 发货状态通用话术库": "shipping_template",
    "16 异常物流话术": "logistics_exception",
    "对标品牌与授权（有部分重复信息）": "brand_comparison",
}


def _headers(row: tuple[Any, ...]) -> list[str]:
    return [normalize_text(value) for value in row]


def _record(headers: list[str], row: tuple[Any, ...]) -> dict[str, str]:
    data: dict[str, str] = {}
    for header, value in zip(headers, row):
        if header:
            data[header] = normalize_text(value)
    return data


def _split_aliases(*values: str) -> list[str]:
    aliases: list[str] = []
    for value in values:
        for part in value.replace("\n", "；").replace("|", "；").replace("、", "；").split("；"):
            text = normalize_text(part)
            if text and text not in aliases:
                aliases.append(text)
    return aliases


def _insert_doc(
    conn,
    *,
    kb_version: str,
    business_type: str,
    product: str,
    topic: str,
    facts: dict[str, str],
    source_sheet: str,
    source_row: int,
    source_field: str = "",
) -> str:
    text_parts = [business_type, product, topic]
    text_parts.extend(f"{key}: {value}" for key, value in facts.items() if value)
    text = "\n".join(part for part in text_parts if part)
    kb_doc_id = stable_id(kb_version, business_type, source_sheet, source_row, product, topic)
    params = (
        kb_doc_id,
        kb_version,
        business_type,
        product,
        topic,
        text,
        json_dumps(facts),
        source_sheet,
        source_row,
        source_field,
    )
    if getattr(conn, "is_pg", False):
        conn.execute(
            """
            INSERT INTO kb_docs
                (kb_doc_id, kb_version, business_type, product, topic, text, facts_json,
                 source_sheet, source_row, source_field)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT (kb_doc_id) DO UPDATE SET
                kb_version = EXCLUDED.kb_version,
                business_type = EXCLUDED.business_type,
                product = EXCLUDED.product,
                topic = EXCLUDED.topic,
                text = EXCLUDED.text,
                facts_json = EXCLUDED.facts_json,
                source_sheet = EXCLUDED.source_sheet,
                source_row = EXCLUDED.source_row,
                source_field = EXCLUDED.source_field
            """,
            params,
        )
    else:
        conn.execute(
            """
            INSERT OR REPLACE INTO kb_docs
                (kb_doc_id, kb_version, business_type, product, topic, text, facts_json,
                 source_sheet, source_row, source_field)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            params,
        )
    return kb_doc_id


def _insert_aliases(conn, aliases: list[str], product: str, kb_doc_id: str) -> None:
    for alias in aliases:
        if getattr(conn, "is_pg", False):
            conn.execute(
                """
                INSERT INTO kb_aliases (alias, product, kb_doc_id, weight)
                VALUES (?, ?, ?, ?)
                ON CONFLICT (alias, product, kb_doc_id) DO NOTHING
                """,
                (alias, product, kb_doc_id, 1.0),
            )
        else:
            conn.execute(
                "INSERT OR IGNORE INTO kb_aliases (alias, product, kb_doc_id, weight) VALUES (?, ?, ?, ?)",
                (alias, product, kb_doc_id, 1.0),
            )


def _first_existing(data: dict[str, str], keys: list[str]) -> str:
    return first_non_empty(data.get(key, "") for key in keys)


def _import_products(conn, ws, kb_version: str) -> int:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return 0
    headers = _headers(rows[0])
    count = 0
    for source_row, row in enumerate(rows[1:], start=2):
        data = _record(headers, row)
        common = data.get("产品常用名", "")
        product = first_non_empty([data.get("产品全称"), common])
        if not product:
            continue
        facts = {key: value for key, value in data.items() if value}
        facts.update(
            {
                "产品常用名": common,
                "产品全称": product,
                "别称": data.get("别称", ""),
                "适用年龄段": data.get("适用年龄段", ""),
                "服用方法": data.get("服用方法（含时间）", ""),
                "使用禁忌": data.get("使用禁忌", ""),
                "产品之间搭配禁忌": data.get("产品之间搭配禁忌", ""),
                "主要成分及含量": data.get("主要成分及含量（每份）", ""),
                "产品规格": data.get("产品规格", ""),
            }
        )
        kb_doc_id = _insert_doc(
            conn,
            kb_version=kb_version,
            business_type="product_profile",
            product=product,
            topic="产品常规信息",
            facts=facts,
            source_sheet=PRODUCT_SHEET,
            source_row=source_row,
        )
        aliases = _split_aliases(common, product, facts["别称"])
        _insert_aliases(conn, aliases, product, kb_doc_id)
        count += 1
    return count


def _import_shipping(conn, ws, kb_version: str) -> int:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return 0
    headers = _headers(rows[0])
    count = 0
    for source_row, row in enumerate(rows[1:], start=2):
        data = _record(headers, row)
        common = data.get("产品常用名", "")
        product = first_non_empty([data.get("产品全称"), common])
        if not product:
            continue
        facts = {
            "产品常用名": common,
            "产品全称": product,
            "发货状态": data.get("发货状态", ""),
            "更新发货时间": data.get("更新发货时间", ""),
            "自定义话术内容": data.get("自定义话术内容", ""),
        }
        kb_doc_id = _insert_doc(
            conn,
            kb_version=kb_version,
            business_type="shipping_status",
            product=product,
            topic="发货状态",
            facts=facts,
            source_sheet=SHIPPING_SHEET,
            source_row=source_row,
        )
        _insert_aliases(conn, _split_aliases(common, product), product, kb_doc_id)
        count += 1
    return count


def _import_recommendations(conn, ws, kb_version: str) -> int:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return 0
    headers = _headers(rows[0])
    count = 0
    for source_row, row in enumerate(rows[1:], start=2):
        data = _record(headers, row)
        product = data.get("推荐产品", "")
        if not product:
            continue
        facts = {
            "需求点": data.get("需求点", ""),
            "挖需结果": data.get("挖需结果", ""),
            "推荐产品": product,
            "推荐产品介绍": data.get("推荐产品介绍", ""),
            "备注": data.get("备注（不发出）", ""),
        }
        kb_doc_id = _insert_doc(
            conn,
            kb_version=kb_version,
            business_type="recommendation_rule",
            product=product,
            topic=facts["需求点"],
            facts=facts,
            source_sheet=RECOMMENDATION_SHEET,
            source_row=source_row,
        )
        _insert_aliases(conn, _split_aliases(product), product, kb_doc_id)
        count += 1
    return count


def _import_reports(conn, ws, kb_version: str) -> int:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return 0
    headers = _headers(rows[0])
    count = 0
    for source_row, row in enumerate(rows[1:], start=2):
        data = _record(headers, row)
        common = first_non_empty([data.get("产品常用名"), data.get("产品")])
        product = first_non_empty([data.get("产品全称"), common])
        if not product:
            continue
        facts = {key: value for key, value in data.items() if value}
        kb_doc_id = _insert_doc(
            conn,
            kb_version=kb_version,
            business_type="quality_compliance",
            product=product,
            topic="产品检测报告",
            facts=facts,
            source_sheet=REPORT_SHEET,
            source_row=source_row,
        )
        _insert_aliases(conn, _split_aliases(common, product), product, kb_doc_id)
        count += 1
    return count


def _import_notices(conn, ws, kb_version: str) -> int:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return 0
    headers = _headers(rows[0])
    count = 0
    for source_row, row in enumerate(rows[1:], start=2):
        data = _record(headers, row)
        title = data.get("通知名称", "")
        speech = data.get("通知话术", "")
        if not (title or speech):
            continue
        facts = {
            "通知名称": title,
            "通知话术": speech,
            "通知开始时间": data.get("通知开始时间", ""),
            "通知结束时间": data.get("通知结束时间", ""),
        }
        kb_doc_id = _insert_doc(
            conn,
            kb_version=kb_version,
            business_type="activity_rule",
            product=title,
            topic="限时通知",
            facts=facts,
            source_sheet=NOTICE_SHEET,
            source_row=source_row,
            source_field="通知话术",
        )
        _insert_aliases(conn, _split_aliases(title), title, kb_doc_id)
        count += 1
    return count


def _import_generic_sheet(conn, ws, kb_version: str, *, source_sheet: str, business_type: str) -> int:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return 0
    headers = _headers(rows[0])
    count = 0
    for source_row, row in enumerate(rows[1:], start=2):
        data = _record(headers, row)
        facts = {key: value for key, value in data.items() if value}
        if not facts:
            continue
        product = _first_existing(
            facts,
            [
                "产品全称",
                "产品常用名",
                "产品",
                "推荐产品",
                "通知名称",
                "标题",
                "问题",
                "原料名称",
                "原料",
                "对标品牌",
                "品牌",
            ],
        )
        topic = _first_existing(
            facts,
            [
                "需求点",
                "通知名称",
                "问题",
                "标题",
                "分类",
                "类别",
                "话术类型",
                "论文方向",
                "功效",
            ],
        )
        if not product:
            product = topic or source_sheet
        if not topic:
            topic = source_sheet
        kb_doc_id = _insert_doc(
            conn,
            kb_version=kb_version,
            business_type=business_type,
            product=product,
            topic=topic,
            facts=facts,
            source_sheet=source_sheet,
            source_row=source_row,
            source_field="row",
        )
        _insert_aliases(conn, _split_aliases(product, topic), product, kb_doc_id)
        count += 1
    return count


def import_workbook(xlsx_path: str | Path, db_path: str | Path, *, kb_version: str = "v1") -> dict[str, int]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    conn = connect(db_path)
    try:
        ensure_schema(conn)
        conn.execute("DELETE FROM kb_meta")
        conn.execute("DELETE FROM kb_docs")
        conn.execute("DELETE FROM kb_aliases")
        counts: dict[str, int] = {}
        if PRODUCT_SHEET in wb.sheetnames:
            counts["product_profile"] = _import_products(conn, wb[PRODUCT_SHEET], kb_version)
        if SHIPPING_SHEET in wb.sheetnames:
            counts["shipping_status"] = _import_shipping(conn, wb[SHIPPING_SHEET], kb_version)
        if RECOMMENDATION_SHEET in wb.sheetnames:
            counts["recommendation_rule"] = _import_recommendations(conn, wb[RECOMMENDATION_SHEET], kb_version)
        if REPORT_SHEET in wb.sheetnames:
            counts["quality_compliance"] = _import_reports(conn, wb[REPORT_SHEET], kb_version)
        if NOTICE_SHEET in wb.sheetnames:
            counts["activity_rule"] = _import_notices(conn, wb[NOTICE_SHEET], kb_version)
        handled_sheets = {PRODUCT_SHEET, SHIPPING_SHEET, RECOMMENDATION_SHEET, REPORT_SHEET, NOTICE_SHEET}
        for sheet_name, business_type in GENERIC_SHEET_TYPES.items():
            if sheet_name in wb.sheetnames and sheet_name not in handled_sheets:
                counts[business_type] = _import_generic_sheet(
                    conn,
                    wb[sheet_name],
                    kb_version,
                    source_sheet=sheet_name,
                    business_type=business_type,
                )
        if getattr(conn, "is_pg", False):
            conn.execute(
                """
                INSERT INTO kb_meta (key, value)
                VALUES (?, ?)
                ON CONFLICT (key) DO UPDATE SET value = EXCLUDED.value
                """,
                ("current_version", kb_version),
            )
        else:
            conn.execute(
                "INSERT OR REPLACE INTO kb_meta (key, value) VALUES (?, ?)",
                ("current_version", kb_version),
            )
        conn.commit()
        return counts
    finally:
        conn.close()
        wb.close()
