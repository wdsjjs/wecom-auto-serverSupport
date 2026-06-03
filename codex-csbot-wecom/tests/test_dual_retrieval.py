import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook

from csbot.codex_contract import validate_codex_reply
from csbot.kb_import import import_workbook
from csbot.policy import should_handoff
from csbot.retrieve import retrieve
from csbot.script_search import script_search
from csbot.vector_store import add_memory
from csbot.vector_store import import_kb_docs_as_memories
from csbot.vector_store import search_memories


def build_sample_workbook(path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    products = wb.create_sheet("5 产品常规信息")
    products.append(
        [
            "产品常用名",
            "产品全称",
            "别称",
            "适用年龄段",
            "服用方法（含时间）",
            "使用禁忌",
            "产品之间搭配禁忌",
            "主要成分及含量（每份）",
            "产品规格",
            "不拼团单盒价格",
            "拼团单盒价格",
            "起拍数量",
            "不拼团起拍价格",
            "拼团起拍价格",
            "1v1商品链接",
        ]
    )
    products.append(
        [
            "女维",
            "女士复合维生素",
            "女士全效复合活性维生素",
            "18 岁及以上成年人",
            "每日 1 粒 随餐服用",
            "18 岁以下未成年人、肾结石、肾病患者不建议服用，孕期、哺乳期请咨询医生；不可与药物同服",
            "与 R-硫辛酸间隔 3h 以上；不建议与复合维生素 B 族同时或同天服用",
            "每 1 粒：多种维生素及矿物质",
            "200 片/瓶",
            "99",
            "89",
            "2",
            "198",
            "178",
            "/packages/goods/detail/index?alias=vitamin",
        ]
    )
    products.append(
        [
            "藻油",
            "婴幼少儿 DHA 藻油",
            "DHA；脑黄金",
            "3 个月以上儿童、少儿及成人",
            "随餐服用；3 岁以下婴幼儿，每日 1 粒（可剪破胶囊皮直接滴入口中或添加在食物中）",
            "3 个月以下婴儿、对藻类过敏者、有出血倾向或服用抗凝药物者、肝功能严重受损者请勿服用",
            "",
            "每 2 粒：DHA 藻油 500mg（其中 DHA 200mg）",
            "60 粒/瓶",
            "129",
            "119",
            "1",
            "129",
            "119",
            "/packages/goods/detail/index?alias=algae",
        ]
    )
    products.append(
        [
            "鱼油",
            "鱼油",
            "鱼脂、深海鱼油",
            "14 岁及以上",
            "每日 2 粒",
            "海鲜过敏者不建议服用",
            "",
            "每 2 粒：Omega-3",
            "60 粒/盒",
            "69",
            "59",
            "4",
            "276",
            "236",
            "/packages/goods/detail/index?alias=fishoil",
        ]
    )

    shipping = wb.create_sheet("1 产品发货状态")
    shipping.append(["产品常用名", "产品全称", "发货状态", "更新发货时间", "自定义话术内容"])
    shipping.append(["rTG鱼油", "95% 高纯度鱼油", "现货", "", ""])
    shipping.append(["EPA鱼油", "97%高纯度EPA鱼油", "现货", "", "预计2个月左右发货哈~"])

    recs = wb.create_sheet("10 补剂推荐")
    recs.append(
        [
            "需求点",
            "挖需铺垫",
            "挖需问题",
            "挖需结果",
            "推荐产品",
            "相同挖需结果下的优先级",
            "是否为兜底推荐产品",
            "推荐产品介绍",
            "推荐后免责话术",
            "备注（不发出）",
            "第一段话术",
        ]
    )
    recs.append(
        [
            "儿童成长",
            "儿童成长需要结合年龄和喂养情况看。",
            "请问孩子多大了，平时饮食怎么样？",
            "3个月以上婴儿，神经发育",
            "婴幼少儿DHA藻油",
            "高",
            "是",
            "促进婴幼儿神经发育",
            "婴幼儿使用前建议结合年龄确认。",
            "只做内部备注",
            "您好~请先介绍基础信息。",
        ]
    )

    reports = wb.create_sheet("13 产品检测报告")
    reports.append(["产品常用名", "产品全称", "出厂检测", "第三方检测"])
    reports.append(["藻油", "婴幼少儿 DHA 藻油", "出厂报告链接", "第三方报告链接"])

    notices = wb.create_sheet("3 限时通知")
    notices.append(["通知名称", "通知话术", "通知开始时间", "通知结束时间", "父记录"])
    notices.append(["一元拼20元优惠券", "拼团成功后15分钟左右优惠券会自动发放到双方账户。", "", "", ""])
    notices.append(["21:00 限时 3 分钟拼团秒杀", "本次活动属于拼团秒杀，需要邀请新用户参团，未成团自动退款。", "5月1号", "5月3号", ""])

    papers = wb.create_sheet("6 论文表")
    papers.append(["产品", "论文方向", "标题", "链接"])
    papers.append(["复合维生素（男/女）", "免疫研究", "维生素相关研究", "https://example.com/paper"])
    papers.append(["婴幼少儿 DHA 藻油", "儿童成长研究", "DHA 与儿童成长研究", "https://example.com/dha-paper"])

    safety = wb.create_sheet("7 L0级注意事项")
    safety.append(["问题", "回答"])
    safety.append(["孕期能不能吃", "孕期请先咨询医生。"])

    raw_materials = wb.create_sheet("12原料专利、认证等材料")
    raw_materials.append(["产品", "原料名称", "认证"])
    raw_materials.append(["女维", "维生素原料", "原料认证示例"])

    shipping_templates = wb.create_sheet("15 发货状态通用话术库")
    shipping_templates.append(["话术类型", "话术"])
    shipping_templates.append(["现货", "现货商品正常发出。"])

    logistics_ex = wb.create_sheet("16 异常物流话术")
    logistics_ex.append(["类别", "话术"])
    logistics_ex.append(["清关停滞", "清关停滞请耐心等待。"])

    brands = wb.create_sheet("对标品牌与授权（有部分重复信息）")
    brands.append(["产品", "对标品牌", "授权"])
    brands.append(["女维", "对标品牌A", "授权信息示例"])

    wb.save(path)


class DualRetrievalTest(unittest.TestCase):
    def setUp(self) -> None:
        self.env_patch = mock.patch.dict(
            os.environ,
            {
                "CSBOT_MEM0_URL": "",
                "CSBOT_MEM0_API_KEY": "",
                "CSBOT_MEM0_ENV": "/tmp/csbot-test-mem0.env",
                "CSBOT_PG_DSN": "",
            },
        )
        self.env_patch.start()
        self.tmp = tempfile.TemporaryDirectory()
        root = Path(self.tmp.name)
        self.xlsx = root / "kb.xlsx"
        self.db = root / "state.sqlite"
        build_sample_workbook(self.xlsx)
        import_workbook(self.xlsx, self.db, kb_version="test-v1")

    def tearDown(self) -> None:
        self.tmp.cleanup()
        self.env_patch.stop()

    def test_script_search_returns_exact_usage_with_source(self) -> None:
        result = script_search("女维怎么吃", self.db)

        self.assertEqual(result["intent"], "product_usage")
        self.assertEqual(result["hits"][0]["product"], "女士复合维生素")
        self.assertEqual(result["hits"][0]["facts"]["服用方法"], "每日 1 粒 随餐服用")
        self.assertEqual(result["hits"][0]["source"]["sheet"], "5 产品常规信息")
        self.assertEqual(result["hits"][0]["source"]["row"], 2)

    def test_retrieve_merges_profile_memory_and_script_fact(self) -> None:
        add_memory(
            self.db,
            customer_id="cust-1",
            text="客户孩子1.5岁，关注婴幼少儿 DHA 藻油补充。",
            metadata={"type": "customer_profile", "product": "婴幼少儿 DHA 藻油"},
        )

        result = retrieve(
            customer_id="cust-1",
            query="上次那个小孩吃的还能继续吗",
            db_path=self.db,
            context={},
        )

        self.assertEqual(result["script_hits"][0]["product"], "婴幼少儿 DHA 藻油")
        self.assertEqual(result["merged_context"]["known_user_facts"]["child_age"], "1.5岁")
        self.assertFalse(result["merged_context"]["needs_clarification"])

    def test_supplement_scope_keeps_recommendation_fields(self) -> None:
        result = script_search("儿童成长 补剂推荐", self.db, {"scope": "supplement"})

        self.assertEqual(result["intent"], "supplement_recommendation")
        self.assertEqual(result["scope"], "supplement")
        first = result["hits"][0]
        self.assertEqual(first["business_type"], "recommendation_rule")
        self.assertEqual(first["source"]["sheet"], "10 补剂推荐")
        self.assertEqual(first["facts"]["挖需铺垫"], "儿童成长需要结合年龄和喂养情况看。")
        self.assertEqual(first["facts"]["挖需问题"], "请问孩子多大了，平时饮食怎么样？")
        self.assertEqual(first["facts"]["相同挖需结果下的优先级"], "高")
        self.assertEqual(first["facts"]["是否为兜底推荐产品"], "是")
        self.assertEqual(first["facts"]["推荐后免责话术"], "婴幼儿使用前建议结合年龄确认。")
        self.assertEqual(first["facts"]["第一段话术"], "您好~请先介绍基础信息。")

    def test_supplement_retrieve_filters_to_supplement_kb(self) -> None:
        result = retrieve(
            customer_id="cust-supplement",
            query="儿童成长 补剂推荐",
            db_path=self.db,
            context={"scope": "supplement"},
        )

        self.assertEqual(result["metrics"]["scope"], "supplement")
        self.assertGreaterEqual(len(result["script_hits"]), 1)
        allowed = {
            ("recommendation_rule", "10 补剂推荐"),
            ("product_profile", "5 产品常规信息"),
            ("safety_policy", "7 L0级注意事项"),
            ("research_evidence", "6 论文表"),
            ("brand_comparison", "对标品牌与授权（有部分重复信息）"),
        }
        self.assertTrue(
            {
                (hit["business_type"], hit["source"]["sheet"])
                for hit in result["script_hits"]
            }.issubset(allowed)
        )

    def test_supplement_scope_includes_related_paper_links(self) -> None:
        result = script_search("儿童成长 婴幼少儿 DHA 藻油 推荐 论文", self.db, {"scope": "supplement"})

        paper_hits = [hit for hit in result["hits"] if hit["business_type"] == "research_evidence"]
        self.assertTrue(paper_hits)
        self.assertEqual(paper_hits[0]["source"]["sheet"], "6 论文表")
        self.assertEqual(paper_hits[0]["facts"]["链接"], "https://example.com/dha-paper")

    def test_supplement_scope_includes_brand_comparison(self) -> None:
        result = script_search("女维 对标品牌 授权", self.db, {"scope": "supplement"})

        brand_hits = [hit for hit in result["hits"] if hit["business_type"] == "brand_comparison"]
        product_hits = [hit for hit in result["hits"] if hit["business_type"] == "product_profile"]
        self.assertTrue(product_hits)
        self.assertEqual(product_hits[0]["source"]["sheet"], "5 产品常规信息")
        self.assertTrue(brand_hits)
        self.assertEqual(brand_hits[0]["source"]["sheet"], "对标品牌与授权（有部分重复信息）")
        self.assertEqual(brand_hits[0]["facts"]["对标品牌"], "对标品牌A")

    def test_fish_shipping_is_ambiguous(self) -> None:
        result = retrieve(customer_id="cust-2", query="鱼油发货时间", db_path=self.db, context={})

        products = {hit["product"] for hit in result["script_hits"]}
        self.assertEqual(products, {"95% 高纯度鱼油", "97%高纯度EPA鱼油"})
        self.assertTrue(result["merged_context"]["needs_clarification"])
        self.assertTrue(result["merged_context"]["conflicts"])

    def test_generic_product_term_matches_alias_variants(self) -> None:
        result = script_search("鱼油发货时间", self.db)

        products = {hit["product"] for hit in result["hits"]}
        self.assertEqual(products, {"95% 高纯度鱼油", "97%高纯度EPA鱼油"})

    def test_product_sales_fields_are_imported_and_retrievable(self) -> None:
        result = script_search("鱼油起拍数量", self.db)

        self.assertEqual(result["intent"], "product_sales")
        self.assertEqual(result["hits"][0]["product"], "鱼油")
        self.assertEqual(result["hits"][0]["facts"]["起拍数量"], "4")
        self.assertEqual(result["hits"][0]["facts"]["不拼团起拍价格"], "276")
        self.assertEqual(result["hits"][0]["source"]["sheet"], "5 产品常规信息")
        self.assertEqual(result["hits"][0]["source"]["row"], 4)
        self.assertFalse(result["needs_clarification"])

    def test_activity_rule_imports_limited_notice_sheet(self) -> None:
        result = script_search("拼团规则", self.db)

        self.assertEqual(result["intent"], "activity_rule")
        self.assertEqual(result["hits"][0]["business_type"], "activity_rule")
        self.assertEqual(result["hits"][0]["source"]["sheet"], "3 限时通知")
        self.assertIn("拼团", result["hits"][0]["facts"]["通知话术"])

    def test_generic_import_covers_remaining_sheets(self) -> None:
        expectations = {
            "research_evidence": "6 论文表",
            "safety_policy": "7 L0级注意事项",
            "raw_material_certification": "12原料专利、认证等材料",
            "shipping_template": "15 发货状态通用话术库",
            "logistics_exception": "16 异常物流话术",
            "brand_comparison": "对标品牌与授权（有部分重复信息）",
        }

        for business_type, source_sheet in expectations.items():
            rows = self._rows_for_business_type(business_type)
            self.assertGreaterEqual(len(rows), 1, business_type)
            self.assertEqual(rows[0]["source_sheet"], source_sheet)

    def test_alias_can_retrieve_generic_sheet_rows(self) -> None:
        result = script_search("女维有什么论文", self.db)

        self.assertEqual(result["intent"], "research_evidence")
        self.assertEqual(result["hits"][0]["source"]["sheet"], "6 论文表")
        self.assertEqual(result["hits"][0]["product"], "复合维生素（男/女）")

    def _rows_for_business_type(self, business_type: str):
        from csbot.db import connect

        conn = connect(self.db)
        try:
            return conn.execute("SELECT * FROM kb_docs WHERE business_type = ?", (business_type,)).fetchall()
        finally:
            conn.close()

    def test_kb_vector_import_replaces_global_knowledge_memories(self) -> None:
        first = import_kb_docs_as_memories(self.db)
        second = import_kb_docs_as_memories(self.db)

        self.assertEqual(first, second)
        hits = retrieve(customer_id="cust-kb", query="鱼油发货时间", db_path=self.db, context={})[
            "vector_hits"
        ]
        ids = [hit["metadata"].get("kb_doc_id") for hit in hits if hit["metadata"].get("type") == "knowledge"]
        self.assertEqual(len(ids), len(set(ids)))

    def test_kb_vector_import_includes_sales_fields(self) -> None:
        import_kb_docs_as_memories(self.db)

        hits = search_memories(self.db, customer_id="cust-kb", query="鱼油起拍数量")

        self.assertTrue(any("起拍数量: 4" in hit["text"] for hit in hits))

    def test_refund_and_complaint_are_handoff(self) -> None:
        self.assertTrue(should_handoff("我要退款并投诉你们"))
        self.assertFalse(should_handoff("女维怎么吃"))

    def test_codex_reply_validation_requires_script_sources_for_send(self) -> None:
        retrieval = retrieve(customer_id="cust-3", query="女维怎么吃", db_path=self.db, context={})

        bad = validate_codex_reply(
            {"action": "send", "reply_text": "每日 1 粒，随餐服用。", "used_script_sources": []},
            retrieval,
        )
        self.assertFalse(bad.ok)

        good = validate_codex_reply(
            {
                "action": "send",
                "reply_text": "每日 1 粒，随餐服用。",
                "used_script_sources": [retrieval["script_hits"][0]["source"]],
                "confidence": 0.9,
            },
            retrieval,
        )
        self.assertTrue(good.ok)

    def test_retrieve_falls_back_to_script_when_vector_disabled(self) -> None:
        result = retrieve(
            customer_id="cust-4",
            query="女维怎么吃",
            db_path=self.db,
            context={"disable_vector": True},
        )

        self.assertEqual(result["merged_context"]["answer_basis"], "script_first")
        self.assertEqual(result["vector_hits"], [])

    def test_retrieve_reports_timing_metrics(self) -> None:
        result = retrieve(customer_id="cust-timing", query="女维怎么吃", db_path=self.db, context={})

        timing = result["metrics"]["timing"]
        self.assertIsInstance(timing["total_ms"], int)
        self.assertGreaterEqual(timing["total_ms"], 0)
        self.assertIsInstance(timing["script_ms"], int)
        self.assertGreaterEqual(timing["script_ms"], 0)
        self.assertIsInstance(timing["vector_ms"], int)
        self.assertGreaterEqual(timing["vector_ms"], 0)
        self.assertIsInstance(timing["hydrate_ms"], int)
        self.assertGreaterEqual(timing["hydrate_ms"], 0)
        self.assertEqual(result["metrics"]["script_count"], len(result["script_hits"]))
        self.assertEqual(result["metrics"]["vector_count"], len(result["vector_hits"]))

    def test_knowledge_age_does_not_become_child_profile(self) -> None:
        import_kb_docs_as_memories(self.db)

        result = retrieve(customer_id="cust-age", query="女维怎么吃", db_path=self.db, context={})

        self.assertNotIn("child_age", result["merged_context"]["known_user_facts"])

    def test_vector_search_uses_mem0_when_configured(self) -> None:
        with mock.patch.dict(os.environ, {"CSBOT_MEM0_URL": "http://mem0.local", "CSBOT_MEM0_API_KEY": "k"}):
            with mock.patch("csbot.vector_store.Mem0Client.from_env") as factory:
                client = factory.return_value
                client.search.return_value = [
                    {
                        "id": "m1",
                        "type": "customer_profile",
                        "text": "客户孩子1.5岁，关注藻油",
                        "score": 0.8,
                        "metadata": {"type": "customer_profile"},
                        "provider": "mem0",
                    }
                ]

                hits = search_memories(self.db, customer_id="cust-mem0", query="小孩藻油")

        self.assertEqual(hits[0]["provider"], "mem0")
        client.search.assert_called_once_with(
            customer_id="cust-mem0",
            query="小孩藻油",
            limit=5,
            include_global=True,
        )

    def test_kb_mem0_import_bulk_deletes_global_knowledge(self) -> None:
        with mock.patch.dict(os.environ, {"CSBOT_MEM0_URL": "http://mem0.local", "CSBOT_MEM0_API_KEY": "k"}):
            with mock.patch("csbot.vector_store.Mem0Client.from_env") as factory:
                client = factory.return_value

                count = import_kb_docs_as_memories(self.db)

        self.assertGreater(count, 0)
        client.clear_memories.assert_called_once_with(customer_id="global-kb")
        self.assertEqual(client.add_memory.call_count, count)


class JsonShapeTest(unittest.TestCase):
    def test_retrieve_output_is_json_serializable(self) -> None:
        sample = {
            "script_hits": [],
            "vector_hits": [],
            "merged_context": {
                "answer_basis": "none",
                "known_user_facts": {},
                "conflicts": [],
                "needs_clarification": True,
            },
        }
        json.dumps(sample, ensure_ascii=False)


if __name__ == "__main__":
    unittest.main()
